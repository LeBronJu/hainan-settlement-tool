from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .monthly_settlement import (
    build_split_files,
    build_summary,
    read_ledger_rows,
    resolve_summary_sheet_name,
    write_report,
)


YEAR = 2026
LEDGER_SHEET = "海南2026年售电结算台账"
FIRST_MONTH_START_COL = 32  # AF
MONTH_BLOCK_WIDTH = 26
BASE_START_COL = 2
BASE_END_COL = 28
POWER_HEADERS = ["零售用户名称", "总电量(I)", "尖段电量(L)", "峰段电量(P)", "平段电量(T)", "谷段电量(X)"]


def s(value: Any) -> str:
    return "" if value is None else str(value).strip()


def n(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def customer_key(value: Any) -> str:
    return "".join(s(value).split())


def ps_string(value: str | Path) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def month_start_col(month: int) -> int:
    if month < 1:
        raise ValueError("月份必须大于等于 1")
    return FIRST_MONTH_START_COL + (month - 1) * MONTH_BLOCK_WIDTH


def main_sheet(wb):
    if LEDGER_SHEET in wb.sheetnames:
        return wb[LEDGER_SHEET]
    if "Sheet1" in wb.sheetnames:
        return wb["Sheet1"]
    for ws in wb.worksheets:
        if "售电结算台账" in s(ws["A1"].value):
            return ws
    raise ValueError(f"找不到台账主表：{LEDGER_SHEET}")


def copy_cell_style(src, dst) -> None:
    import copy

    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy.copy(src.alignment)
    if src.font:
        dst.font = copy.copy(src.font)
    if src.fill:
        dst.fill = copy.copy(src.fill)
    if src.border:
        dst.border = copy.copy(src.border)
    if src.protection:
        dst.protection = copy.copy(src.protection)
    if src.comment:
        dst.comment = copy.copy(src.comment)


def translated_value(src, target_row: int, target_col: int) -> Any:
    value = src.value
    if isinstance(value, str) and value.startswith("="):
        try:
            return Translator(value, origin=src.coordinate).translate_formula(
                row_delta=target_row - src.row,
                col_delta=target_col - src.column,
            )
        except Exception:
            return value
    return value


def load_json_from_powershell(ps: str, timeout: int = 180) -> Any:
    completed = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=timeout,
    )
    text = completed.stdout.strip()
    if not text:
        return []
    data = json.loads(text)
    return data if isinstance(data, list) else [data]


def read_raw_power_rows_with_excel(raw_path: Path) -> list[dict[str, Any]]:
    ps = f"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$path = {ps_string(raw_path)}
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
$items = New-Object System.Collections.Generic.List[object]
try {{
  $wb = $excel.Workbooks.Open($path)
  $ws = $wb.Worksheets.Item(1)
  $rows = $ws.UsedRange.Rows.Count
  for ($r = 4; $r -le $rows; $r++) {{
    $name = [string]$ws.Cells.Item($r, 4).Text
    if ($name.Trim()) {{
      $items.Add([pscustomobject]@{{
        name = $name.Trim()
        total = $ws.Cells.Item($r, 9).Value2
        sharp = $ws.Cells.Item($r, 12).Value2
        peak = $ws.Cells.Item($r, 16).Value2
        flat = $ws.Cells.Item($r, 20).Value2
        valley = $ws.Cells.Item($r, 24).Value2
      }}) | Out-Null
    }}
  }}
  $wb.Close($false)
}} finally {{
  $excel.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
$items | ConvertTo-Json -Depth 4
"""
    return load_json_from_powershell(ps)


def read_raw_power_rows_from_xlsx(raw_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(raw_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows: list[dict[str, Any]] = []
    for row in range(4, ws.max_row + 1):
        name = ws.cell(row, 4).value
        if not s(name):
            continue
        rows.append(
            {
                "name": s(name),
                "total": ws.cell(row, 9).value,
                "sharp": ws.cell(row, 12).value,
                "peak": ws.cell(row, 16).value,
                "flat": ws.cell(row, 20).value,
                "valley": ws.cell(row, 24).value,
            }
        )
    return rows


def read_raw_power_rows_from_csv(raw_path: Path) -> list[dict[str, Any]]:
    encodings = ["utf-8-sig", "gbk"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            with raw_path.open("r", encoding=encoding, newline="") as fh:
                reader = csv.reader(fh)
                rows = list(reader)[3:]
            break
        except Exception as exc:
            last_error = exc
    else:
        raise last_error or ValueError(f"无法读取 CSV：{raw_path}")

    result: list[dict[str, Any]] = []
    for row in rows:
        if len(row) <= 23 or not s(row[3]):
            continue
        result.append(
            {
                "name": s(row[3]),
                "total": row[8],
                "sharp": row[11],
                "peak": row[15],
                "flat": row[19],
                "valley": row[23],
            }
        )
    return result


def read_raw_power_rows(raw_path: Path) -> list[dict[str, Any]]:
    suffix = raw_path.suffix.lower()
    if suffix == ".csv":
        return read_raw_power_rows_from_csv(raw_path)
    if suffix == ".xlsx":
        return read_raw_power_rows_from_xlsx(raw_path)
    if suffix == ".xls":
        return read_raw_power_rows_with_excel(raw_path)
    raise ValueError(f"暂不支持的原始明细格式：{raw_path.suffix}")


def write_power_workbook(rows: list[dict[str, Any]], output_path: Path) -> dict[str, Any]:
    grouped: dict[str, dict[str, Any]] = {}
    display_name: dict[str, str] = {}
    for item in rows:
        key = customer_key(item["name"])
        if not key:
            continue
        display_name.setdefault(key, s(item["name"]))
        current = grouped.setdefault(
            key,
            {"total": 0.0, "sharp": 0.0, "peak": 0.0, "flat": 0.0, "valley": 0.0},
        )
        current["total"] += n(item.get("total"))
        current["sharp"] += n(item.get("sharp"))
        current["peak"] += n(item.get("peak"))
        current["flat"] += n(item.get("flat"))
        current["valley"] += n(item.get("valley"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(POWER_HEADERS)
    for key in sorted(grouped, key=lambda value: display_name[value]):
        values = grouped[key]
        ws.append(
            [
                display_name[key],
                round(values["total"], 4),
                round(values["sharp"], 4),
                round(values["peak"], 4),
                round(values["flat"], 4),
                round(values["valley"], 4),
            ]
        )
    for col, width in enumerate([36, 14, 14, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "output": str(output_path),
        "raw_rows": len(rows),
        "customers": len(grouped),
        "total_power": round(sum(value["total"] for value in grouped.values()), 4),
    }


def clean_power_data(raw_path: Path, output_path: Path) -> dict[str, Any]:
    if not raw_path.exists():
        raise FileNotFoundError(f"找不到原始明细：{raw_path}")
    rows = read_raw_power_rows(raw_path)
    return write_power_workbook(rows, output_path)


def read_power_rows(power_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(power_path, data_only=True)
    ws = wb.active
    rows: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 1).value
        if not s(name):
            continue
        rows.append(
            {
                "source_row": row,
                "name": s(name),
                "key": customer_key(name),
                "total": ws.cell(row, 2).value or 0,
                "sharp": ws.cell(row, 3).value or 0,
                "peak": ws.cell(row, 4).value or 0,
                "flat": ws.cell(row, 5).value or 0,
                "valley": ws.cell(row, 6).value or 0,
            }
        )
    return rows


def read_customer_codes_with_excel(raw_path: Path) -> dict[str, str]:
    ps = f"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$path = {ps_string(raw_path)}
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
$map = @{{}}
try {{
  $wb = $excel.Workbooks.Open($path)
  foreach ($sheetName in @('零售主体电量', '零售户号电量')) {{
    try {{ $ws = $wb.Worksheets.Item($sheetName) }} catch {{ continue }}
    $rows = $ws.UsedRange.Rows.Count
    for ($r = 4; $r -le $rows; $r++) {{
      $code = [string]$ws.Cells.Item($r, 3).Text
      $name = [string]$ws.Cells.Item($r, 4).Text
      if ($name.Trim() -and $code.Trim() -and -not $map.ContainsKey($name.Trim())) {{
        $map[$name.Trim()] = $code.Trim()
      }}
    }}
  }}
  $wb.Close($false)
}} finally {{
  $excel.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
$map | ConvertTo-Json -Depth 4
"""
    try:
        data = load_json_from_powershell(ps)
    except Exception:
        return {}
    if not data:
        return {}
    if isinstance(data[0], dict) and "name" not in data[0]:
        raw = data[0]
        return {customer_key(name): s(code) for name, code in raw.items() if s(name) and s(code)}
    return {}


def read_customer_codes_from_xlsx(raw_path: Path) -> dict[str, str]:
    wb = load_workbook(raw_path, data_only=True, read_only=True)
    result: dict[str, str] = {}
    sheet_names = [name for name in ("零售主体电量", "零售户号电量") if name in wb.sheetnames]
    if not sheet_names:
        sheet_names = [wb.sheetnames[0]]
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for row in range(4, ws.max_row + 1):
            code = s(ws.cell(row, 3).value)
            name = s(ws.cell(row, 4).value)
            if name and code:
                result.setdefault(customer_key(name), code)
    return result


def read_customer_codes(raw_path: Path | None) -> dict[str, str]:
    if not raw_path or not raw_path.exists():
        return {}
    if raw_path.suffix.lower() == ".xlsx":
        return read_customer_codes_from_xlsx(raw_path)
    if raw_path.suffix.lower() == ".xls":
        return read_customer_codes_with_excel(raw_path)
    return {}


def ledger_row_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in range(4, ws.max_row + 1):
        seq = ws.cell(row, 1).value
        name = ws.cell(row, 3).value
        if isinstance(seq, int) and s(name):
            result[customer_key(name)] = row
    return result


def unmerge_overlapping_block(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for merged in list(ws.merged_cells.ranges):
        if (
            merged.min_row <= max_row
            and merged.max_row >= min_row
            and merged.min_col <= max_col
            and merged.max_col >= min_col
        ):
            ws.unmerge_cells(str(merged))


def copy_month_block(ws, source_month: int, target_month: int) -> None:
    source_start = month_start_col(source_month)
    target_start = month_start_col(target_month)
    col_offset = target_start - source_start
    max_row = ws.max_row
    target_end = target_start + MONTH_BLOCK_WIDTH - 1

    # If the target month was generated before, its old merged cells make
    # non-anchor cells read-only. Clear only the target month block first.
    unmerge_overlapping_block(ws, 1, max_row, target_start, target_end)

    for col in range(source_start, source_start + MONTH_BLOCK_WIDTH):
        src_letter = get_column_letter(col)
        dst_letter = get_column_letter(col + col_offset)
        ws.column_dimensions[dst_letter].width = ws.column_dimensions[src_letter].width
        ws.column_dimensions[dst_letter].hidden = ws.column_dimensions[src_letter].hidden
        ws.column_dimensions[dst_letter].outlineLevel = ws.column_dimensions[src_letter].outlineLevel

    for row in range(1, max_row + 1):
        for col in range(source_start, source_start + MONTH_BLOCK_WIDTH):
            src = ws.cell(row, col)
            dst = ws.cell(row, col + col_offset)
            copy_cell_style(src, dst)
            dst.value = translated_value(src, row, col + col_offset)

    existing = {str(rng) for rng in ws.merged_cells.ranges}
    for merged in list(ws.merged_cells.ranges):
        if merged.min_col >= source_start and merged.max_col < source_start + MONTH_BLOCK_WIDTH:
            new_range = (
                f"{get_column_letter(merged.min_col + col_offset)}{merged.min_row}:"
                f"{get_column_letter(merged.max_col + col_offset)}{merged.max_row}"
            )
            if new_range not in existing:
                ws.merge_cells(new_range)
                existing.add(new_range)
    ws.cell(1, target_start).value = f"{target_month}月"


def copy_row_template(ws, style_template_row: int, formula_template_row: int, target_row: int, target_month: int) -> None:
    target_start = month_start_col(target_month)
    max_col = target_start + MONTH_BLOCK_WIDTH
    ws.row_dimensions[target_row].height = ws.row_dimensions[style_template_row].height
    for col in range(1, max_col):
        src = ws.cell(style_template_row, col)
        dst = ws.cell(target_row, col)
        copy_cell_style(src, dst)
        formula_src = ws.cell(formula_template_row, col)
        if (
            isinstance(formula_src.value, str)
            and formula_src.value.startswith("=")
            and (col in {23, 24, 25, 27} or col >= target_start + 5)
        ):
            dst.value = translated_value(formula_src, target_row, col)
        else:
            dst.value = None


def copy_base_info_from_reference(out_ws, ref_ws, out_row: int, ref_row: int) -> None:
    for col in range(BASE_START_COL, BASE_END_COL + 1):
        src = ref_ws.cell(ref_row, col)
        dst = out_ws.cell(out_row, col)
        copy_cell_style(src, dst)
        dst.value = translated_value(src, out_row, col)


def set_recalc(wb) -> None:
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def excel_recalculate(paths: list[Path], timeout: int = 240) -> str:
    real_paths = [path for path in dict.fromkeys(paths) if path.exists() and not path.name.startswith("~$")]
    if not real_paths:
        return "No Excel files to recalculate"
    path_lines = ",\n".join(f"  {ps_string(path)}" for path in real_paths)
    ps = f"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$paths = @(
{path_lines}
)
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
  foreach ($p in $paths) {{
    $wb = $excel.Workbooks.Open($p)
    $excel.CalculateFullRebuild()
    $wb.Save()
    $wb.Close($true)
  }}
  Write-Output 'Excel recalculation saved'
}} finally {{
  $excel.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
"""
    completed = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=timeout,
    )
    return completed.stdout.strip() or "Excel recalculation saved"


def default_ledger_output_name(base_ledger: Path, month: int) -> str:
    stem = base_ledger.stem.rstrip("】")
    return f"{stem}】补{month}月电量.xlsx"


def update_ledger(
    *,
    month: int,
    base_ledger: Path,
    power_path: Path,
    output_dir: Path,
    raw_detail: Path | None = None,
    reference_ledger: Path | None = None,
    output_name: str | None = None,
    copy_reference_existing: bool = False,
) -> tuple[Path, Path, dict[str, Any]]:
    if month <= 1:
        raise ValueError("更新台账要求月份大于 1")
    if not base_ledger.exists():
        raise FileNotFoundError(f"找不到基础台账：{base_ledger}")
    if not power_path.exists():
        raise FileNotFoundError(f"找不到电量表：{power_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / (output_name or default_ledger_output_name(base_ledger, month))
    report_path = output_dir / f"{month}月台账更新报告.json"
    shutil.copy2(base_ledger, out_path)

    out_wb = load_workbook(out_path, data_only=False)
    out_ws = main_sheet(out_wb)
    power_rows = read_power_rows(power_path)
    raw_code_map = read_customer_codes(raw_detail)
    out_map = ledger_row_map(out_ws)

    ref_ws = None
    ref_map: dict[str, int] = {}
    if reference_ledger and reference_ledger.exists():
        ref_wb = load_workbook(reference_ledger, data_only=False)
        ref_ws = main_sheet(ref_wb)
        ref_map = ledger_row_map(ref_ws)

    last_data_row = max(out_map.values())
    next_row = last_data_row + 1
    next_seq = max(out_ws.cell(row, 1).value for row in out_map.values()) + 1

    target_start = month_start_col(month)
    target_month_already_present = s(out_ws.cell(1, target_start).value) == f"{month}月"
    existing_target_month_values = []
    for row in out_map.values():
        value = out_ws.cell(row, target_start).value
        if n(value) > 0:
            existing_target_month_values.append(
                {"row": row, "name": s(out_ws.cell(row, 3).value), "total": value}
            )

    if not target_month_already_present:
        copy_month_block(out_ws, month - 1, month)

    matched: list[dict[str, Any]] = []
    new_rows: list[dict[str, Any]] = []
    copied_from_reference: list[dict[str, Any]] = []
    missing_reference: list[str] = []
    code_filled_from_raw: list[dict[str, Any]] = []
    missing_codes: list[str] = []

    for item in power_rows:
        found_existing = item["key"] in out_map
        target_row = out_map.get(item["key"])
        if target_row:
            matched.append({"name": item["name"], "target_row": target_row, "total": item["total"]})
        else:
            target_row = next_row
            copy_row_template(out_ws, last_data_row, 4, target_row, month)
            out_ws.cell(target_row, 1).value = next_seq
            out_ws.cell(target_row, 3).value = item["name"]
            out_map[item["key"]] = target_row
            new_rows.append({"name": item["name"], "target_row": target_row, "total": item["total"]})
            next_row += 1
            next_seq += 1

        ref_row = ref_map.get(item["key"])
        should_copy_ref = ref_ws is not None and ref_row and ((not found_existing) or copy_reference_existing)
        if should_copy_ref:
            copy_base_info_from_reference(out_ws, ref_ws, target_row, ref_row)
            out_ws.cell(target_row, 1).value = out_ws.cell(target_row, 1).value or target_row - 3
            out_ws.cell(target_row, 3).value = item["name"]
            copied_from_reference.append({"name": item["name"], "target_row": target_row, "reference_row": ref_row})
        elif not found_existing and ref_ws is not None:
            missing_reference.append(item["name"])

        if not s(out_ws.cell(target_row, 2).value):
            code = raw_code_map.get(item["key"])
            if code:
                out_ws.cell(target_row, 2).value = code
                code_filled_from_raw.append({"name": item["name"], "target_row": target_row, "code": code})
            else:
                missing_codes.append(item["name"])

        out_ws.cell(target_row, target_start + 0).value = item["total"]
        out_ws.cell(target_row, target_start + 1).value = item["sharp"]
        out_ws.cell(target_row, target_start + 2).value = item["peak"]
        out_ws.cell(target_row, target_start + 3).value = item["flat"]
        out_ws.cell(target_row, target_start + 4).value = item["valley"]

    set_recalc(out_wb)
    out_wb.save(out_path)

    duplicate_power = {key: count for key, count in Counter(item["key"] for item in power_rows).items() if count > 1}
    missing_manual = []
    for item in power_rows:
        row = out_map[item["key"]]
        if not s(out_ws.cell(row, 10).value):
            missing_manual.append({"row": row, "name": item["name"], "missing": "负责人/J列"})

    report = {
        "month": month,
        "source_ledger": str(base_ledger),
        "source_power": str(power_path),
        "raw_detail_for_codes": str(raw_detail) if raw_detail else None,
        "reference_ledger": str(reference_ledger) if reference_ledger else None,
        "output": str(out_path),
        "target_block": f"{get_column_letter(target_start)}:{get_column_letter(target_start + MONTH_BLOCK_WIDTH - 1)}",
        "target_month_already_present": target_month_already_present,
        "power_rows": len(power_rows),
        "matched_rows": len(matched),
        "new_rows": len(new_rows),
        "new_customers": new_rows,
        "existing_target_month_values": existing_target_month_values,
        "copied_from_reference": copied_from_reference,
        "missing_reference": missing_reference,
        "raw_detail_code_rows": len(raw_code_map),
        "code_filled_from_raw": code_filled_from_raw,
        "missing_codes": missing_codes,
        "missing_manual_info": missing_manual,
        "duplicate_names_in_power_file": duplicate_power,
        "month_total": round(sum(float(item["total"] or 0) for item in power_rows), 4),
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path, report_path, report


def validate_ledger_for_month(ledger_path: Path, month: int) -> dict[str, Any]:
    wb = load_workbook(ledger_path, data_only=True)
    ws = main_sheet(wb)
    start = month_start_col(month)
    missing_codes = []
    missing_owner = []
    active_rows = 0
    for row in range(4, ws.max_row + 1):
        name = ws.cell(row, 3).value
        total = ws.cell(row, start).value
        if not s(name) or n(total) <= 0:
            continue
        active_rows += 1
        if not s(ws.cell(row, 2).value):
            missing_codes.append({"row": row, "name": s(name)})
        if not s(ws.cell(row, 10).value):
            missing_owner.append({"row": row, "name": s(name)})
    return {
        "ledger": str(ledger_path),
        "month": month,
        "active_rows": active_rows,
        "missing_codes": missing_codes,
        "missing_owner": missing_owner,
    }


def run_settlement(
    *,
    month: int,
    ledger_path: Path,
    template_root: Path,
    summary_template: Path,
    output_dir: Path,
    final_summary_name: str | None = None,
    allow_missing_owner: bool = False,
    recalc_excel: bool = True,
) -> tuple[Path, Path, dict[str, Any]]:
    if not ledger_path.exists():
        raise FileNotFoundError(f"找不到台账：{ledger_path}")
    if not template_root.exists():
        raise FileNotFoundError(f"找不到分表模板文件夹：{template_root}")
    if not summary_template.exists():
        raise FileNotFoundError(f"找不到汇总表模板：{summary_template}")
    output_dir.mkdir(parents=True, exist_ok=True)

    ledger_check = validate_ledger_for_month(ledger_path, month)
    if ledger_check["missing_owner"] and not allow_missing_owner:
        names = "、".join(item["name"] for item in ledger_check["missing_owner"][:10])
        raise ValueError(f"{month}月台账还有负责人缺失，先不要生成分表/汇总表：{names}")

    proxy_rows, inter_rows = read_ledger_rows(ledger_path, month)
    owner_missing_rows = [
        {"ledger_row": row.ledger_row, "customer": row.customer, "entity": row.entity, "kind": row.kind}
        for row in [*proxy_rows, *inter_rows]
        if not s(row.owner)
    ]
    if owner_missing_rows and not allow_missing_owner:
        names = "、".join(item["customer"] for item in owner_missing_rows[:10])
        raise ValueError(f"{month}月结算明细存在负责人缺失：{names}")

    totals = build_split_files(template_root, output_dir, month, proxy_rows, inter_rows)
    summary_path = build_summary(summary_template, output_dir, month, totals)
    final_summary_path = output_dir / (final_summary_name or f"【2026年海南省代理费汇总表-{month}月自动化】.xlsx")
    if summary_path != final_summary_path:
        shutil.copy2(summary_path, final_summary_path)
        if summary_path.exists():
            summary_path.unlink()

    settlement_report_path = write_report(output_dir, month, totals, proxy_rows, inter_rows, final_summary_path)
    recalc_status = None
    if recalc_excel:
        recalc_paths = [ledger_path, final_summary_path]
        recalc_paths.extend(Path(item.output_file) for item in totals)
        recalc_status = excel_recalculate(recalc_paths)

    summary_check = validate_summary(final_summary_path, month)
    report = {
        "month": month,
        "ledger": str(ledger_path),
        "template_root": str(template_root),
        "summary_template": str(summary_template),
        "output_dir": str(output_dir),
        "summary": str(final_summary_path),
        "settlement_report": str(settlement_report_path),
        "ledger_check": ledger_check,
        "summary_check": summary_check,
        "proxy_rows": len(proxy_rows),
        "intermediary_rows": len(inter_rows),
        "proxy_groups": sum(1 for item in totals if item.kind == "代理费"),
        "intermediary_groups": sum(1 for item in totals if item.kind == "居间费"),
        "proxy_total": round(sum(item.expected_net for item in totals if item.kind == "代理费"), 4),
        "intermediary_total": round(sum(item.expected_net for item in totals if item.kind == "居间费"), 4),
        "excel_recalculation": recalc_status,
    }
    report_path = output_dir / f"{month}月结算生成总报告.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return final_summary_path, report_path, report


def validate_summary(summary_path: Path, month: int) -> dict[str, Any]:
    wb = load_workbook(summary_path, data_only=True)
    result: dict[str, Any] = {"summary": str(summary_path), "sheets": {}}
    for role in ("main", "qingneng", "qinghui"):
        sheet_name = resolve_summary_sheet_name(wb, role, required=(role == "main"))
        if not sheet_name:
            continue
        ws = wb[sheet_name]
        headers = [s(ws.cell(2, col).value) for col in range(1, ws.max_column + 1)]
        result["sheets"][sheet_name] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "last_column": get_column_letter(ws.max_column),
            "has_month_header": f"{YEAR}年{month}月" in headers,
            "too_many_columns": ws.max_column > 200,
        }
    return result


def find_default_raw(month_root: Path) -> Path | None:
    candidates = []
    for suffix in (".xls", ".xlsx", ".csv"):
        candidates.extend(month_root.glob(f"零售侧明细结果*{suffix}"))
    clean_words = ("数据处理", "数据清洗")
    filtered = [path for path in candidates if not any(word in path.name for word in clean_words)]
    return sorted(filtered, key=lambda path: path.stat().st_mtime, reverse=True)[0] if filtered else None


def run_all(args: argparse.Namespace) -> dict[str, Any]:
    month_root = Path(args.month_root)
    output_dir = Path(args.output_dir) if args.output_dir else month_root / f"{args.month}月自动化输出"
    output_dir.mkdir(parents=True, exist_ok=True)

    raw_path = Path(args.raw_detail) if args.raw_detail else find_default_raw(month_root)
    power_path = Path(args.power) if args.power else month_root / "零售侧用户电量数据处理表.xlsx"
    clean_report = None
    if not power_path.exists():
        if not raw_path:
            raise FileNotFoundError("找不到电量表，也没有找到可清洗的零售侧明细结果文件")
        clean_report = clean_power_data(raw_path, power_path)

    ledger_path, ledger_report_path, ledger_report = update_ledger(
        month=args.month,
        base_ledger=Path(args.base_ledger),
        power_path=power_path,
        output_dir=output_dir,
        raw_detail=raw_path,
        reference_ledger=Path(args.reference_ledger) if args.reference_ledger else None,
        output_name=args.ledger_name,
        copy_reference_existing=args.copy_reference_existing,
    )

    summary_path, settlement_report_path, settlement_report = run_settlement(
        month=args.month,
        ledger_path=ledger_path,
        template_root=Path(args.template_root),
        summary_template=Path(args.summary_template),
        output_dir=output_dir,
        final_summary_name=args.summary_name,
        allow_missing_owner=args.allow_missing_owner,
        recalc_excel=not args.no_excel_recalc,
    )

    final = {
        "month": args.month,
        "output_dir": str(output_dir),
        "raw_detail": str(raw_path) if raw_path else None,
        "power": str(power_path),
        "ledger": str(ledger_path),
        "summary": str(summary_path),
        "clean_report": clean_report,
        "ledger_report": str(ledger_report_path),
        "settlement_report": str(settlement_report_path),
        "ledger_generation": ledger_report,
        "settlement_generation": settlement_report,
    }
    final_path = output_dir / f"{args.month}月自动化流程报告.json"
    final_path.write_text(json.dumps(final, ensure_ascii=False, indent=2), encoding="utf-8")
    return final


def add_common_month(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--month", type=int, required=True, help="月份，例如 5")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="海南售电月度结算自动化流程")
    sub = parser.add_subparsers(dest="command", required=True)

    clean = sub.add_parser("clean", help="清洗零售侧明细结果，生成零售侧用户电量数据处理表")
    clean.add_argument("--raw-detail", required=True, help="原始零售侧明细结果 .xls/.xlsx/.csv")
    clean.add_argument("--output", required=True, help="输出的电量表 .xlsx")

    ledger = sub.add_parser("ledger", help="把当月电量更新进台账")
    add_common_month(ledger)
    ledger.add_argument("--base-ledger", required=True, help="上月或修正版台账")
    ledger.add_argument("--power", required=True, help="当月电量表")
    ledger.add_argument("--output-dir", required=True, help="输出文件夹")
    ledger.add_argument("--raw-detail", help="原始明细，用于补企业编号")
    ledger.add_argument("--reference-ledger", help="已人工补齐/确认后的本月台账，用于给新增客户复制基础资料")
    ledger.add_argument("--ledger-name", help="输出台账文件名")
    ledger.add_argument("--copy-reference-existing", action="store_true", help="也用参考台账覆盖已有客户的基础资料，默认只复制新增客户")

    settlement = sub.add_parser("settlement", help="由已补齐台账生成代理/居间分表和汇总表")
    add_common_month(settlement)
    settlement.add_argument("--ledger", required=True, help="已补齐的本月台账")
    settlement.add_argument("--template-root", required=True, help="上月结算文件夹，里面要有 2026年代理 - 海南 / 2026年居间 - 海南")
    settlement.add_argument("--summary-template", required=True, help="上月或修正版代理费汇总表")
    settlement.add_argument("--output-dir", required=True, help="输出文件夹")
    settlement.add_argument("--summary-name", help="输出汇总表文件名")
    settlement.add_argument("--allow-missing-owner", action="store_true", help="允许负责人缺失时继续生成，默认不允许")
    settlement.add_argument("--no-excel-recalc", action="store_true", help="不调用 Excel COM 重算")

    all_cmd = sub.add_parser("all", help="一键跑：必要时清洗电量 -> 更新台账 -> 生成分表/汇总表")
    add_common_month(all_cmd)
    all_cmd.add_argument("--month-root", required=True, help="当月结算文件夹")
    all_cmd.add_argument("--base-ledger", required=True, help="上月或修正版台账")
    all_cmd.add_argument("--template-root", required=True, help="上月结算文件夹/分表模板根目录")
    all_cmd.add_argument("--summary-template", required=True, help="上月或修正版汇总表")
    all_cmd.add_argument("--raw-detail", help="原始零售侧明细结果；不填则从当月文件夹自动找")
    all_cmd.add_argument("--power", help="当月电量表；不填则使用/生成 零售侧用户电量数据处理表.xlsx")
    all_cmd.add_argument("--reference-ledger", help="已人工补齐/确认后的本月台账，用于新增客户基础资料")
    all_cmd.add_argument("--output-dir", help="输出文件夹；默认：当月文件夹/<月份>月自动化输出")
    all_cmd.add_argument("--ledger-name", help="输出台账文件名")
    all_cmd.add_argument("--summary-name", help="输出汇总表文件名")
    all_cmd.add_argument("--copy-reference-existing", action="store_true", help="也用参考台账覆盖已有客户的基础资料，默认只复制新增客户")
    all_cmd.add_argument("--allow-missing-owner", action="store_true", help="允许负责人缺失时继续生成，默认不允许")
    all_cmd.add_argument("--no-excel-recalc", action="store_true", help="不调用 Excel COM 重算")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.command == "clean":
        result = clean_power_data(Path(args.raw_detail), Path(args.output))
    elif args.command == "ledger":
        ledger_path, report_path, report = update_ledger(
            month=args.month,
            base_ledger=Path(args.base_ledger),
            power_path=Path(args.power),
            output_dir=Path(args.output_dir),
            raw_detail=Path(args.raw_detail) if args.raw_detail else None,
            reference_ledger=Path(args.reference_ledger) if args.reference_ledger else None,
            output_name=args.ledger_name,
            copy_reference_existing=args.copy_reference_existing,
        )
        result = {"ledger": str(ledger_path), "report": str(report_path), **report}
    elif args.command == "settlement":
        summary_path, report_path, report = run_settlement(
            month=args.month,
            ledger_path=Path(args.ledger),
            template_root=Path(args.template_root),
            summary_template=Path(args.summary_template),
            output_dir=Path(args.output_dir),
            final_summary_name=args.summary_name,
            allow_missing_owner=args.allow_missing_owner,
            recalc_excel=not args.no_excel_recalc,
        )
        result = {"summary": str(summary_path), "report": str(report_path), **report}
    elif args.command == "all":
        result = run_all(args)
    else:
        raise ValueError(f"未知命令：{args.command}")

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
